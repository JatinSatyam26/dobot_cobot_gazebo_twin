#!/usr/bin/env python3
# Copyright (c) 2026 Jatin Satyam
# SPDX-License-Identifier: Apache-2.0
"""Build the seven-questions form as a .docx that can be answered by typing."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INK      = RGBColor(0x16, 0x16, 0x1A)
MUTED    = RGBColor(0x5A, 0x5A, 0x66)
ACCENT   = RGBColor(0xA8, 0x15, 0x4A)
BOX_FILL = "F6F6F9"
BOX_LINE = "D3D3DC"
ACC_LINE = "A8154A"

# OOXML enforces child-element ORDER inside pPr / tblPr / tcPr / rPr. Appending
# is a schema violation Word may reject, so each element is inserted before the
# first of its legal successors (validated with the docx skill's validate.py).
PPR_AFTER_PBDR = ('w:shd', 'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap',
                  'w:overflowPunct', 'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN',
                  'w:bidi', 'w:adjustRightInd', 'w:snapToGrid', 'w:spacing', 'w:ind',
                  'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap', 'w:jc',
                  'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap', 'w:outlineLvl',
                  'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange')
TBLPR_AFTER_BORDERS = ('w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook', 'w:tblCaption',
                       'w:tblDescription', 'w:tblPrChange')
TCPR_AFTER_SHD = ('w:noWrap', 'w:tcMar', 'w:textDirection', 'w:tcFitText', 'w:vAlign',
                  'w:hideMark', 'w:headers', 'w:tcPrChange')
RPR_AFTER_SPACING = ('w:w', 'w:kern', 'w:position', 'w:sz', 'w:szCs', 'w:highlight', 'w:u',
                     'w:effect', 'w:bdr', 'w:shd', 'w:fitText', 'w:vertAlign', 'w:rtl', 'w:cs',
                     'w:em', 'w:lang', 'w:eastAsianLayout', 'w:specVanish', 'w:oMath')

QUESTIONS = [
    (1, False,
     "Is there a network switch at the bench that the PLC and both robots plug into? "
     "Or is each thing wired on its own?",
     None,
     "My laptop has to sit on the same network as them to read anything at all."),
    (2, False,
     "When you program the PLC, do you plug in an Ethernet cable or a USB cable?",
     None,
     "If it's always USB, the PLC's network side may never have been switched on. That's not a "
     "problem — it just becomes a small setup job, and it's better to know now than on the day."),
    (3, True,
     "What turns the vacuum on — the PLC, or the Pro 600 robot's own outputs?",
     "Changes how I build it.",
     "My code assumes the PLC controls it. If the robot does, I've been planning to read the "
     "wrong device."),
    (4, True,
     "Do the PLC and the two robots talk to each other over the network, or are they joined by "
     "plain wires — one sends a “go” signal and the other reacts?",
     "Changes how I build it.",
     "It decides whether the PLC knows what the robots are doing, or whether I have to read all "
     "three separately and line them up myself."),
    (5, False,
     "Who starts a cycle — somebody pressing a button, the PLC on its own, or a person "
     "hitting play on a robot pendant?",
     None,
     "Tells me what signal to watch for as “a new cycle has begun”, so the simulation "
     "starts at the same moment the machine does."),
    (6, True,
     "Are there any sensors on the cell at all — something that detects the holder arriving, "
     "or that a wafer is present?",
     "Changes how I build it.",
     "Right now I assume there are none and work out where the holder is by counting seconds from "
     "when the belt starts. A real sensor would be far more accurate, and I'd rather read it than "
     "guess."),
    (7, False,
     "Does the belt ever run backwards, or does the holder always travel one way and get carried "
     "back by hand?",
     None,
     "In my simulation the belt returns the holder by itself. If the real one doesn't, the two "
     "stop matching after the first cycle."),
]


def set_table_borders(table, color=BOX_LINE, sz=6, left_color=None):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(18 if (edge == "left" and left_color) else sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), left_color if (edge == "left" and left_color) else color)
        borders.append(el)
    tblPr.insert_element_before(borders, *TBLPR_AFTER_BORDERS)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.insert_element_before(shd, *TCPR_AFTER_SHD)


def para(doc, before=0, after=3, spacing=1.0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = spacing
    return p


def run(p, text, size=10.5, bold=False, italic=False, color=INK, caps=False, spacing_pts=None):
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.all_caps = caps
    if spacing_pts is not None:
        rPr = r._element.get_or_add_rPr()
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(int(spacing_pts * 20)))
        rPr.insert_element_before(sp, *RPR_AFTER_SPACING)
    return r


def bottom_rule(p, color="16161A", sz=12):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    bdr.append(bottom)
    pPr.insert_element_before(bdr, *PPR_AFTER_PBDR)


def answer_box(doc, height_cm=1.5, key=False):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    t.columns[0].width = Cm(17.0)
    cell = t.cell(0, 0)
    cell.width = Cm(17.0)
    set_table_borders(t, left_color=ACC_LINE if key else None)
    shade(cell, BOX_FILL)
    row = t.rows[0]
    row.height = Cm(height_cm)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    inner = cell.paragraphs[0]
    inner.paragraph_format.space_after = Pt(0)
    r = inner.add_run("")
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    return t


def build(path):
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(3)

    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0)
    s.top_margin = Cm(1.8)
    s.bottom_margin = Cm(1.6)

    # masthead
    p = para(doc, after=1)
    run(p, "Wafer cell · digital shadow", size=8, bold=True, color=ACCENT, caps=True, spacing_pts=0.6)

    p = para(doc, after=4)
    run(p, "Seven questions about the cell", size=21, bold=True)

    p = para(doc, after=8)
    run(p, "I'm building a simulation that mirrors the real cell on screen while it runs. It ", size=10)
    run(p, "only reads", size=10, bold=True)
    run(p, " — it never sends a command to the PLC or to either robot. These seven answers tell me "
           "where to read from and what the readings mean. Type your answers in the grey boxes. "
           "Short answers are fine; “I don't know” is a useful answer too.", size=10)
    bottom_rule(p)

    # who answered
    t = doc.add_table(rows=1, cols=2)
    t.autofit = False
    widths = (Cm(10.0), Cm(7.0))
    for i, w in enumerate(widths):
        t.columns[i].width = w
        t.cell(0, i).width = w
    set_table_borders(t)
    for i, label in enumerate(("Answered by", "Date")):
        c = t.cell(0, i)
        shade(c, BOX_FILL)
        cp = c.paragraphs[0]
        cp.paragraph_format.space_after = Pt(0)
        run(cp, label + ":  ", size=8.5, bold=True, color=MUTED, caps=True, spacing_pts=0.5)
    t.rows[0].height = Cm(0.85)
    t.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    p = para(doc, before=10, after=7)
    run(p, "■  ", size=10, color=ACCENT)
    run(p, "The three marked in this colour change how I build it, so they're the ones I'd like first.",
        size=9, color=MUTED)

    for num, key, question, flag, why in QUESTIONS:
        p = para(doc, before=6, after=1)
        run(p, f"{num}.  ", size=11.5, bold=True, color=ACCENT if key else INK)
        run(p, question, size=11.5, bold=True)

        p = para(doc, after=3)
        if flag:
            run(p, flag + " ", size=9, bold=True, color=ACCENT)
        run(p, why, size=9, color=MUTED)

        answer_box(doc, key=key)

    p = para(doc, before=12, after=0)
    bottom_rule(p, color="D3D3DC", sz=6)
    p = para(doc, before=4, after=0)
    run(p, "Read-only: nothing on this sheet changes how the machine runs.", size=8.5, color=MUTED)

    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")

    doc.save(path)
    return path


if __name__ == "__main__":
    import sys
    print("wrote", build(sys.argv[1]))
