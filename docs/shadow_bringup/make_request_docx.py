#!/usr/bin/env python3
# Copyright (c) 2026 Jatin Satyam
# SPDX-License-Identifier: Apache-2.0
"""Build the full 'Cell Shadow Data Request' as a .docx, ending with the seven questions."""
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn

from make_seven_questions_docx import (INK, MUTED, ACCENT, BOX_FILL, BOX_LINE, ACC_LINE,
                       set_table_borders, shade, para, run, bottom_rule, QUESTIONS)

SOFT_FILL = "FBEAF0"      # accent-tinted, for the "say it like this" boxes
STEEL     = RGBColor(0x2E, 0x5C, 0x7E)
AMBER     = RGBColor(0x85, 0x56, 0x0A)
HEAD_FILL = "16161A"
ALT_FILL  = "F2F2F5"

CHIP_COLOR = {"BLOCKER": ACCENT, "HELPFUL": STEEL, "MAY NEED A CHANGE": AMBER, "BEST CASE": STEEL}


# ----------------------------------------------------------------- building blocks
def part_head(doc, label, title):
    p = para(doc, before=16, after=2)
    run(p, label, size=8.5, bold=True, color=ACCENT, caps=True, spacing_pts=0.7)
    p = para(doc, after=3)
    run(p, title, size=15.5, bold=True)
    return p


def lede(doc, text):
    p = para(doc, after=8)
    run(p, text, size=10, color=MUTED)
    return p


def ask_head(doc, title, chips):
    p = para(doc, before=11, after=2)
    for c in chips:
        run(p, c + "   ", size=8, bold=True, color=CHIP_COLOR[c], caps=True, spacing_pts=0.6)
    p2 = para(doc, after=2)
    run(p2, title, size=12, bold=True)
    return p2


def say_box(doc, text):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    t.columns[0].width = Cm(17.0)
    cell = t.cell(0, 0)
    cell.width = Cm(17.0)
    set_table_borders(t, color=SOFT_FILL, left_color=ACC_LINE)
    shade(cell, SOFT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run(p, "SAY IT LIKE THIS", size=7.5, bold=True, color=ACCENT, caps=True, spacing_pts=0.7)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    run(p2, text, size=10.5, italic=True)
    para(doc, after=0)
    return t


def why(doc, lead, text):
    p = para(doc, before=4, after=2)
    if lead:
        run(p, lead + " ", size=9.5, bold=True, color=INK)
    run(p, text, size=9.5, color=MUTED)
    return p


def where(doc, label, items):
    p = para(doc, before=5, after=2)
    run(p, label, size=7.5, bold=True, color=MUTED, caps=True, spacing_pts=0.7)
    for i, item in enumerate(items, 1):
        p = para(doc, after=1)
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.7)
        run(p, f"{i}.  ", size=9.5, bold=True, color=MUTED)
        run(p, item, size=9.5, color=MUTED)


def callout(doc, label, text, color=AMBER, fill="F8EEDA"):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    t.columns[0].width = Cm(17.0)
    cell = t.cell(0, 0)
    cell.width = Cm(17.0)
    set_table_borders(t, color=fill, left_color="85560A")
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run(p, label, size=7.5, bold=True, color=color, caps=True, spacing_pts=0.7)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    run(p2, text, size=10, color=INK)
    para(doc, after=0)


def answer_box(doc, height_cm=1.5, key=False, width_cm=17.0):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    t.columns[0].width = Cm(width_cm)
    cell = t.cell(0, 0)
    cell.width = Cm(width_cm)
    set_table_borders(t, left_color=ACC_LINE if key else None)
    shade(cell, BOX_FILL)
    row = t.rows[0]
    row.height = Cm(height_cm)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    inner = cell.paragraphs[0]
    inner.paragraph_format.space_after = Pt(0)
    r = inner.add_run("")
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    para(doc, after=0)
    return t


# ----------------------------------------------------------------- content
MESSAGE = [
    ("p", "Hi — I'm building a simulation of our wafer cell that mirrors the real machine on "
          "screen while it runs. It only reads. It never sends a command to the PLC or either robot."),
    ("p", "To hook it up I need a few things from your side. None of it is urgent, and none of it "
          "changes your program except possibly item 3."),
    ("n", "The exact PLC model. It's printed on the front of the controller — something like "
          "2080-LC50-24QWB. A photo of the label is perfect."),
    ("n", "Its IP address and subnet mask. Also: is the PLC plugged into a network switch, or do "
          "you only connect a USB cable when you program it?"),
    ("n", "The Global Variables list from your CCW project. Open the project, find Global Variables "
          "in the tree on the left, and send me a screenshot of the table — I need the names and the "
          "data types. One important thing: if the belt and vacuum flags only exist as local "
          "variables inside the program, I can't read them over the network. If that's the case, "
          "could you add Global Variables that mirror them?"),
    ("n", "Which variable is which. Which one turns the belt on, which one sets the belt direction, "
          "which one turns the vacuum on, and is there one that holds a step or state number for the "
          "cycle? Exact spelling matters."),
    ("n", "Does the belt ever run backwards, and roughly how many seconds does it run in each "
          "direction per cycle?"),
    ("n", "Permission: is it OK if I plug my laptop into the same switch and read values while the "
          "cell is running? Read only, nothing written."),
    ("p", "If it's easier, just send me the whole CCW project archive (the .ccwarc file) and I'll "
          "pull the details out myself."),
]

ASKS_PLC = [
    (["BLOCKER"], "1.1  Which controller is it, exactly?",
     "“What's the exact model number on the front of the PLC? Something like 2080-LC50-24QWB. "
     "A photo of the label works.”",
     [("Why it comes first:", "not every controller in this family has an Ethernet port. If this one "
       "doesn't, everything else on this page changes, and we'd need a different way to get the "
       "signals out."),
      (None, "The number also tells me how many inputs and outputs it has, which is a sanity check "
       "against the signal list in 1.4.")],
     None),
    (["BLOCKER"], "1.2  What is its network address?",
     "“What's the PLC's IP address and subnet mask? And when you program it, do you plug in an "
     "Ethernet cable or a USB cable?”",
     [("If he says USB:", "that's a real possibility, and it means the Ethernet port may never have "
       "been configured. It's not a problem — it just turns this from a question into a small setup "
       "task for him, and it's better to find out now than on demo day.")],
     ("Where to look", [
         "Some controllers show the address on a small display on the front.",
         "Otherwise it's in the CCW project, in the controller's settings under the Ethernet page.",
         "Menu names shift a little between CCW versions — a screenshot of whatever he finds is fine."])),
    (["BLOCKER", "MAY NEED A CHANGE"], "1.3  The Global Variables list",
     "“Can you open Global Variables in the CCW project and send me a screenshot of the table — the "
     "names and the data types? And can you check that the belt and vacuum flags are Global, not "
     "local to the program?”",
     [("This is the one item that may need him to edit his program.",
       "On this family of controller, only Global Variables can be read from the network. A variable "
       "that lives inside a program is invisible from outside, no matter what I do at my end."),
      (None, "If his flags are local, the fix is small and safe: add a Global Variable for each one "
       "and assign the local value to it in one rung. It doesn't change how the machine behaves.")],
     ("Where to look", [
         "Open the project in Connected Components Workbench.",
         "In the tree on the left, under the controller, there's a Global Variables entry.",
         "Open it — it's a table with Name, Data Type, Initial Value and Comment.",
         "A screenshot is enough. The whole project archive is better."])),
    (["BLOCKER"], "1.4  Which variable does what",
     "“For each of these, what's the exact variable name and its type — belt on/off, belt direction, "
     "vacuum on/off, and a step or state number if the program keeps one?”",
     [(None, "Spelling and capitals have to be exact, the way they're written in the table. My side "
       "asks for a variable by name; a near miss just returns nothing."),
      ("The step number is the valuable one.",
       "If his program tracks where it is in the cycle as a number, the simulation can follow the "
       "real machine step by step. Without it I have to guess the sequence from the belt and vacuum "
       "turning on and off — that works, but it's guesswork."),
      ("What I'm currently guessing:",
       "the names in my config today are placeholders — BeltRun, VacuumOn, CycleStep. Almost "
       "certainly wrong. Correcting a wrong guess is faster than answering a blank question, so it's "
       "worth showing him this line.")],
     None),
    (["HELPFUL"], "1.5  A look at the program itself",
     "“Could you send a screenshot of the main ladder routine, or the whole project archive? I want "
     "to see the order the cycle runs in.”",
     [(None, "The simulation has its own list of cycle steps that I built from the videos. Seeing his "
       "logic lets me line the two up so they use the same names and the same order — that's the "
       "difference between a simulation that looks right and one that genuinely tracks the machine.")],
     None),
    (["HELPFUL"], "1.6  Timings and direction",
     "“How many seconds does the belt run each time? Does it ever go backwards? Is there a pause "
     "partway along? And are there any sensors on the belt, or is it all on timers?”",
     [(None, "There are no belt sensors as far as I can tell, so my simulation works out where the "
       "holder is by counting time from when the belt starts. If that's right, I need his timer "
       "values. If there is a sensor, that's much better and I'd rather read it.")],
     None),
    (["BLOCKER"], "1.7  Permission to connect",
     "“Is it alright if I plug my laptop into the same switch and read values while the cell is "
     "running? Read only — I won't write anything or go online with the controller.”",
     [(None, "Worth asking plainly rather than assuming. A PLC programmer's instinct when someone "
       "connects to their controller is that the program is about to be changed. Saying “read only, "
       "I'm not going online with it” up front removes that worry.")],
     None),
]

ASKS_ROBOTS = [
    (["BLOCKER"], "2.1  The addresses of both arms",
     "“What IP addresses do the Dobot M1 Pro and the myCobot Pro 600 have? They're usually shown on "
     "the pendant's network or settings screen.”",
     [(None, "My current guesses are the vendor defaults — 192.168.1.6 for the Dobot and "
       "192.168.1.159 for the Pro 600. Real benches rarely keep those."),
      (None, "For the Pro 600 there's a second part: its remote-control server has to be switched on "
       "for anything to read from it. Ask whether it's enabled, and if he doesn't know, that's fine — "
       "we can find out by trying.")],
     None),
    (["BLOCKER"], "2.2  The one-joint-at-a-time test",
     "“Can we jog one joint at a time and write down what the pendant says? Move joint 1 on its own, "
     "note the number, then joint 2, and so on. Ten minutes for both arms.”",
     [("This is the least obvious ask and one of the most valuable.",
       "The real robot and my model may count the same joint in opposite directions, or from a "
       "different zero point. Without this the simulated arm mirrors the real one but bends the "
       "wrong way."),
      (None, "It's a five-minute job per arm and it removes a whole class of confusing errors later. "
       "Doing it together is much faster than describing it over messages.")],
     None),
    (["HELPFUL"], "2.3  The parked position",
     "“When both arms are sitting at their start-of-cycle position, can you photograph the joint "
     "numbers on each pendant?”",
     [(None, "My rest poses were calculated, not measured from the real machine. Real numbers would "
       "replace a guess I've been carrying since the beginning.")],
     None),
]

ASKS_NET = [
    (["BLOCKER"], "3.1  Can I get on the same network?",
     "“Is there a network switch at the bench with a spare port? What range are the addresses in — "
     "do they all start 192.168.1.something? And is anything else on that network, like a company "
     "router?”",
     [(None, "If the three devices sit on an isolated switch, I plug in, give my laptop an address in "
       "the same range, and everything works. If they're on the building network, there may be rules "
       "about what can be plugged in — better to know before carrying a laptop over.")],
     None),
]

SHEET = [
    ("PLC model", "Micro850 family"),
    ("PLC address", "192.168.1.20"),
    ("Belt on/off variable", "BeltRun (BOOL)"),
    ("Belt direction variable", "none — belt never reverses"),
    ("Vacuum variable", "VacuumOn (BOOL)"),
    ("Step / state number", "CycleStep — may not exist"),
    ("Global or local?", "assumed Global"),
    ("Belt speed", "0.07 m/s, measured off a video"),
    ("Belt run time per move", "unknown"),
    ("Pause partway along the belt", "none seen in the video"),
    ("Belt sensors", "none — timers only"),
    ("Dobot M1 Pro address", "192.168.1.6"),
    ("Pro 600 address", "192.168.1.159, port 5001"),
    ("Pro 600 remote server on?", "assumed yes"),
    ("Joint directions, both arms", "assumed same as my model"),
]


def render_asks(doc, asks):
    for chips, title, say, whys, wh in asks:
        ask_head(doc, title, chips)
        say_box(doc, say)
        for lead, text in whys:
            why(doc, lead, text)
        if wh:
            where(doc, wh[0], wh[1])


def build(path):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(3)

    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0)
    s.top_margin = Cm(1.8)
    s.bottom_margin = Cm(1.6)

    # ---------- masthead
    p = para(doc, after=1)
    run(p, "Wafer cell · digital shadow · bring-up", size=8, bold=True, color=ACCENT,
        caps=True, spacing_pts=0.7)
    p = para(doc, after=4)
    run(p, "What I need from the PLC side", size=21, bold=True)
    p = para(doc, after=6)
    run(p, "A read-only list. Everything here is either a number to look up or a screenshot to "
           "send — except one item, which may need a small change to the PLC program.", size=10.5,
        color=MUTED)
    p = para(doc, after=8)
    run(p, "Cell: Dobot M1 Pro + myCobot Pro 600 + conveyor      PLC: Micro800 family, "
           "programmed in CCW", size=8.5, color=MUTED)
    bottom_rule(p)

    # ---------- who answered
    t = doc.add_table(rows=1, cols=2)
    t.autofit = False
    for i, w in enumerate((Cm(10.0), Cm(7.0))):
        t.columns[i].width = w
        t.cell(0, i).width = w
    set_table_borders(t)
    for i, label in enumerate(("Answered by", "Date")):
        c = t.cell(0, i)
        shade(c, BOX_FILL)
        cp = c.paragraphs[0]
        cp.paragraph_format.space_after = Pt(0)
        run(cp, label + ":  ", size=8.5, bold=True, color=MUTED, caps=True, spacing_pts=0.5)
    t.rows[0].height = Cm(0.85)
    t.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    para(doc, after=0)

    # ---------- the message
    part_head(doc, "Start here", "The message to send him")
    lede(doc, "If you only send one thing, send this. Everything after it is backup detail for when "
              "he asks “what do you mean?”")
    tm = doc.add_table(rows=1, cols=1)
    tm.autofit = False
    tm.columns[0].width = Cm(17.0)
    mc = tm.cell(0, 0)
    mc.width = Cm(17.0)
    set_table_borders(tm)
    shade(mc, "FAFAFC")
    first = True
    n = 0
    for kind, text in MESSAGE:
        mp = mc.paragraphs[0] if first else mc.add_paragraph()
        first = False
        mp.paragraph_format.space_after = Pt(5)
        if kind == "n":
            n += 1
            mp.paragraph_format.left_indent = Cm(0.75)
            mp.paragraph_format.first_line_indent = Cm(-0.75)
            run(mp, f"{n}.  ", size=10, bold=True)
            run(mp, text, size=10)
        else:
            run(mp, text, size=10)
    para(doc, after=0)

    # ---------- parts
    part_head(doc, "Part 1", "The PLC — his home turf")
    lede(doc, "Take these in order. Each one only makes sense once the one before it is answered.")
    render_asks(doc, ASKS_PLC)

    part_head(doc, "Part 2", "The two robots")
    lede(doc, "He may not own these, but he'll know who does, or the numbers may be on the teach "
              "pendants.")
    render_asks(doc, ASKS_ROBOTS)

    part_head(doc, "Part 3", "The bench network")
    lede(doc, "Small, quick, and it decides whether any of the above can actually be read on the day.")
    render_asks(doc, ASKS_NET)

    part_head(doc, "Part 4", "Twenty minutes together beats twenty messages")
    lede(doc, "If he's willing to sit at the bench with you once, this replaces most of the "
              "back-and-forth.")
    where(doc, "What we'd do in that session", [
        "Plug my laptop into the switch and confirm I can see all three devices.",
        "Read one variable from the PLC while he watches — proves the whole path works and "
        "reassures him nothing is being written.",
        "Run one cycle while I record everything the three devices report.",
        "Do the one-joint-at-a-time test on both arms."])
    why(doc, None, "That recording alone would let me finish the shadow without him being present "
                   "again.")

    # ---------- answer sheet
    part_head(doc, "Sheet", "Answer sheet")
    lede(doc, "Fill in the right-hand column. The middle column is what I'm assuming today — most of "
              "it is probably wrong, and that's fine.")
    widths = (Cm(5.4), Cm(6.6), Cm(5.0))
    ts = doc.add_table(rows=len(SHEET) + 1, cols=3)
    ts.autofit = False
    set_table_borders(ts)
    for i, w in enumerate(widths):
        ts.columns[i].width = w
    for i, head in enumerate(("What", "My guess today", "Actual")):
        c = ts.cell(0, i)
        c.width = widths[i]
        shade(c, HEAD_FILL)
        cp = c.paragraphs[0]
        cp.paragraph_format.space_after = Pt(0)
        run(cp, head, size=8.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), caps=True,
            spacing_pts=0.6)
    for r_i, (what, guess) in enumerate(SHEET, start=1):
        for c_i, text in enumerate((what, guess, "")):
            c = ts.cell(r_i, c_i)
            c.width = widths[c_i]
            if c_i == 2:
                shade(c, BOX_FILL)
            elif r_i % 2 == 0:
                shade(c, ALT_FILL)
            cp = c.paragraphs[0]
            cp.paragraph_format.space_after = Pt(0)
            if text:
                run(cp, text, size=9.5, color=MUTED if c_i == 1 else INK)
    para(doc, after=6)
    callout(doc, "If he can only do one thing",
            "Ask for the CCW project archive — the .ccwarc file. It contains the variable list, the "
            "data types and the ladder logic in one attachment, and it answers items 1.3, 1.4, 1.5 "
            "and most of 1.6 at once.")

    # ---------- the seven questions
    part_head(doc, "Questions", "Seven questions about the cell")
    lede(doc, "These are the ones I'd like answered directly, in plain words. Type into the grey "
              "boxes. Short answers are fine; “I don't know” is a useful answer too.")
    p = para(doc, after=7)
    run(p, "■  ", size=10, color=ACCENT)
    run(p, "The three marked in this colour change how I build it, so they're the ones I'd like "
           "first.", size=9, color=MUTED)

    for num, key, question, flag, reason in QUESTIONS:
        p = para(doc, before=7, after=1)
        run(p, f"{num}.  ", size=11.5, bold=True, color=ACCENT if key else INK)
        run(p, question, size=11.5, bold=True)
        p = para(doc, after=3)
        if flag:
            run(p, flag + " ", size=9, bold=True, color=ACCENT)
        run(p, reason, size=9, color=MUTED)
        answer_box(doc, key=key)

    p = para(doc, before=14, after=0)
    bottom_rule(p, color=BOX_LINE, sz=6)
    p = para(doc, before=4, after=0)
    run(p, "Read-only throughout: the simulation observes the cell and never commands it. Nothing "
           "on this list changes how the machine runs, with the single exception of item 1.3, which "
           "may add mirrored Global Variables to the PLC program.", size=8.5, color=MUTED)

    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")

    doc.save(path)
    return path


if __name__ == "__main__":
    print("wrote", build(sys.argv[1]))
